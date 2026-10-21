"""Build browser-ready chapter data from guided notes Word answer sheets.

Usage: python build_guided_notes.py "Chapter 1 Guided Notes.docx" [more chapters...]
"""

import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document


def merge_text(parts, value):
    if not value:
        return
    if parts and parts[-1]["type"] == "text":
        parts[-1]["value"] += value
    else:
        parts.append({"type": "text", "value": value})


def paragraph_parts(paragraph, next_id):
    parts = []
    groups = []
    for run in paragraph.runs:
        if not run.text:
            continue
        is_answer = run.bold is True
        if groups and groups[-1][0] == is_answer:
            groups[-1] = (is_answer, groups[-1][1] + run.text)
        else:
            groups.append((is_answer, run.text))

    for is_answer, value in groups:
        if not is_answer or not value.strip():
            merge_text(parts, value)
            continue
        leading = value[: len(value) - len(value.lstrip())]
        trailing = value[len(value.rstrip()):]
        merge_text(parts, leading)
        parts.append({"type": "blank", "id": f"b{next_id}", "answer": value.strip()})
        next_id += 1
        merge_text(parts, trailing)
    return parts, next_id


def build_chapter(path):
    doc = Document(path)
    if doc.tables:
        raise ValueError(f"{path.name}: tables need an explicit content review")
    title = next((p.text.strip() for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip()), path.stem)
    chapter_match = re.search(r"Chapter\s+(\d+)", title, re.I)
    if not chapter_match:
        raise ValueError(f"{path.name}: no chapter number in first heading")
    number = int(chapter_match.group(1))
    blocks = []
    next_id = 1
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip():
            continue
        style = paragraph.style.name
        if style == "Heading 1":
            continue
        if style == "Heading 2":
            kind = "h2"
        elif style == "Heading 3":
            kind = "h3"
        elif style == "List Paragraph":
            kind = "list"
        else:
            kind = "paragraph"
        parts, next_id = paragraph_parts(paragraph, next_id)
        reconstructed = "".join(part.get("value", part.get("answer", "")) for part in parts)
        if reconstructed != paragraph.text:
            raise ValueError(f"{path.name}: paragraph text changed during extraction: {paragraph.text!r}")
        level = 0
        if kind == "list":
            ppr = paragraph._p.pPr
            if ppr is not None and ppr.numPr is not None and ppr.numPr.ilvl is not None:
                level = int(ppr.numPr.ilvl.val)
        blocks.append({"type": kind, "level": level, "parts": parts})
    blank_count = next_id - 1
    if not blank_count:
        raise ValueError(f"{path.name}: no bold answer text found")
    return {
        "id": f"chapter-{number}",
        "number": number,
        "title": title,
        "source": path.name,
        "sourceSha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "blankCount": blank_count,
        "blocks": blocks,
    }


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    chapters = sorted((build_chapter(Path(arg)) for arg in sys.argv[1:]), key=lambda ch: ch["number"])
    numbers = [chapter["number"] for chapter in chapters]
    if len(numbers) != len(set(numbers)):
        raise ValueError("Duplicate chapter numbers")
    output = Path(__file__).with_name("chapters.js")
    output.write_text(
        "// Generated from the supplied guided notes answer sheets.\n"
        + "window.GUIDED_NOTES_CHAPTERS = "
        + json.dumps(chapters, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )
    for chapter in chapters:
        print(f"Chapter {chapter['number']}: {len(chapter['blocks'])} blocks, {chapter['blankCount']} blanks")
    print(output)


if __name__ == "__main__":
    main()
